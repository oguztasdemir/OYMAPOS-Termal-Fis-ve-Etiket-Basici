# -*- coding: utf-8 -*-
"""
📦 Product Controller
Ürün arama, CRUD, fiyat geçmişi (audit trail), etiket fiyat eşitleme, Excel/CSV dışa aktarma
"""
import os
import io
import csv
from typing import Optional
import tempfile
from fastapi import APIRouter, Form, UploadFile, File
from fastapi.responses import StreamingResponse, JSONResponse

from backend.models.schemas import ProductUpdateRequest
from backend.services.db_service import (
    get_all_products, search_products, get_product_by_barcode, get_products_count,
    sync_all_label_prices_to_pos_price, get_full_product_history, revert_product_history,
    update_product_details, update_products_by_clipboard_data, update_product_printed_time,
    preview_clipboard_price_update
)
from backend.services.vegawin_service import parse_raw_text_products, parse_vegawin_file
from backend.utils.response_utils import success_response, error_response

from backend.services.db.connection import db_session
from backend.utils.text_utils import get_blacklist_data, save_blacklist_data, clean_barcode_text
from pydantic import BaseModel

router = APIRouter(prefix="/api", tags=["Products"])

@router.get("/products")
async def get_products(q: str = "", only_new: bool = False, only_diff: bool = False, only_blacklist: bool = False, limit: int = 0):
    effective_limit = limit if limit > 0 else None
    bl_data = get_blacklist_data()
    blacklist_barcodes = set(str(b).strip() for b in bl_data.get("barcodes", []) if str(b).strip())

    if q.strip():
        items = search_products(q.strip(), limit=effective_limit, only_new=only_new, only_diff=only_diff, only_blacklist=only_blacklist, blacklist_barcodes=blacklist_barcodes)
    else:
        items = get_all_products(limit=effective_limit, only_new=only_new, only_diff=only_diff, only_blacklist=only_blacklist, blacklist_barcodes=blacklist_barcodes)
    
    # Her ürüne kara listede olup olmadığını ekle
    for item in items:
        b_code = str(item.get("barcode", "")).strip()
        item["is_blacklisted"] = (b_code in blacklist_barcodes)

    # İlgili tüm sayaçları hesapla
    counts = {
        "total": get_products_count(blacklist_barcodes=blacklist_barcodes),
        "diff": get_products_count(only_diff=True, blacklist_barcodes=blacklist_barcodes),
        "new": get_products_count(only_new=True, blacklist_barcodes=blacklist_barcodes),
        "blacklist": get_products_count(only_blacklist=True, blacklist_barcodes=blacklist_barcodes)
    }

    return success_response(
        data={
            "products": items,
            "total": get_products_count(only_new=only_new, only_diff=only_diff, only_blacklist=only_blacklist, blacklist_barcodes=blacklist_barcodes),
            "counts": counts,
            "blacklist_barcodes": list(blacklist_barcodes)
        },
        message="Ürünler listelendi"
    )

@router.post("/products/{barcode}/toggle-blacklist")
async def toggle_product_blacklist(barcode: str):
    """Ürünü kara listeye ekler veya kara listeden çıkarır."""
    b = clean_barcode_text(barcode)
    if not b:
        return error_response(message="Geçersiz barkod.", status_code=400)
        
    bl_data = get_blacklist_data()
    barcodes = [str(x).strip() for x in bl_data.get("barcodes", []) if str(x).strip()]
    
    is_now_blacklisted = False
    if b in barcodes:
        barcodes = [x for x in barcodes if x != b]
        bl_data["barcodes"] = barcodes
        save_blacklist_data(bl_data)
        is_now_blacklisted = False
        msg = f"'{b}' barkodlu ürün kara listeden kaldırıldı."
    else:
        barcodes.append(b)
        bl_data["barcodes"] = barcodes
        save_blacklist_data(bl_data)
        is_now_blacklisted = True
        msg = f"'{b}' barkodlu ürün kara listeye eklendi."

    with db_session() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE urunler SET is_blacklisted = ? WHERE barcode = ?;", (1 if is_now_blacklisted else 0, b))
        conn.commit()

    b_set = set(barcodes)
    counts = {
        "total": get_products_count(blacklist_barcodes=b_set),
        "diff": get_products_count(only_diff=True, blacklist_barcodes=b_set),
        "new": get_products_count(only_new=True, blacklist_barcodes=b_set),
        "blacklist": get_products_count(only_blacklist=True, blacklist_barcodes=b_set)
    }

    return success_response(
        data={
            "barcode": b,
            "is_blacklisted": is_now_blacklisted,
            "counts": counts
        },
        message=msg
    )

class BatchBlacklistRequest(BaseModel):
    barcodes: list
    action: str  # 'add' veya 'remove'

@router.post("/products/batch-blacklist")
async def batch_blacklist_products(req: BatchBlacklistRequest):
    """Seçilen birden fazla ürünü topluca kara listeye ekler veya çıkarır."""
    raw_barcodes = [clean_barcode_text(b) for b in req.barcodes if clean_barcode_text(b)]
    if not raw_barcodes:
        return error_response(message="İşlem yapılacak ürün seçilmedi.", status_code=400)

    bl_data = get_blacklist_data()
    current_barcodes = set(str(x).strip() for x in bl_data.get("barcodes", []) if str(x).strip())
    
    is_add = (req.action == 'add')
    if is_add:
        for b in raw_barcodes:
            current_barcodes.add(b)
    else:
        for b in raw_barcodes:
            current_barcodes.discard(b)

    bl_data["barcodes"] = list(current_barcodes)
    save_blacklist_data(bl_data)

    with db_session() as conn:
        cursor = conn.cursor()
        val = 1 if is_add else 0
        cursor.executemany("UPDATE urunler SET is_blacklisted = ? WHERE barcode = ?;", [(val, b) for b in raw_barcodes])
        conn.commit()

    counts = {
        "total": get_products_count(blacklist_barcodes=current_barcodes),
        "diff": get_products_count(only_diff=True, blacklist_barcodes=current_barcodes),
        "new": get_products_count(only_new=True, blacklist_barcodes=current_barcodes),
        "blacklist": get_products_count(only_blacklist=True, blacklist_barcodes=current_barcodes)
    }

    action_text = "kara listeye eklendi" if is_add else "kara listeden çıkarıldı"
    return success_response(
        data={
            "updated_count": len(raw_barcodes),
            "action": req.action,
            "counts": counts
        },
        message=f"{len(raw_barcodes)} ürün başarıyla {action_text}."
    )

class TogglePriceChangeRequest(BaseModel):
    title: Optional[str] = None
    price: Optional[float] = None
    device_name: Optional[str] = "Mobil Reyon Terminali"

@router.get("/products/{barcode}")
async def get_product(barcode: str):
    prod = get_product_by_barcode(barcode)
    if not prod:
        return error_response(message="Ürün bulunamadı.", status_code=404)
    
    # Bugünün değişenlerinde var mı kontrol et
    today_str = datetime.datetime.now().strftime("%Y-%m-%d")
    with db_session() as conn:
        c = conn.cursor()
        c.execute("""
            SELECT 1 FROM product_history 
            WHERE barcode = ? AND substr(created_at, 1, 10) = ? AND sync_id IS NULL 
            LIMIT 1;
        """, (clean_barcode_text(barcode), today_str))
        prod["is_in_today_changes"] = (c.fetchone() is not None)

    return success_response(data={"product": prod}, message="Ürün bulundu")

@router.post("/products/{barcode}/toggle-price-change")
async def toggle_product_price_change(barcode: str, req: Optional[TogglePriceChangeRequest] = None):
    """
    Ürünü bugünkü fiyat değişimleri raporuna manuel ekler veya varsa kaldırır (Toggle).
    """
    b = clean_barcode_text(barcode)
    if not b:
        return error_response(message="Geçersiz barkod.", status_code=400)
    
    today_str = datetime.datetime.now().strftime("%Y-%m-%d")
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    dev_name = req.device_name if req and req.device_name else "Mobil Reyon Terminali"

    with db_session() as conn:
        cursor = conn.cursor()
        
        # 1. Ürün veritabanında var mı?
        cursor.execute("SELECT * FROM urunler WHERE barcode = ? LIMIT 1;", (b,))
        prod = cursor.fetchone()

        if not prod:
            new_title = req.title.strip() if req and req.title and req.title.strip() else f"Yeni Ürün ({b})"
            new_price = float(req.price) if req and req.price is not None else 0.0
            cursor.execute("""
                INSERT INTO urunler (
                    barcode, stock_code, title, raw_system_title, price, label_price,
                    brand, unit, source_device, is_new, is_blacklisted,
                    created_at, updated_at, price_updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (
                b, b, new_title, new_title, new_price, 0.0,
                "", "ADET", dev_name, 1, 0,
                now_str, now_str, now_str
            ))
            cur_title = new_title
            cur_price = new_price
            label_p = 0.0
        else:
            p_dict = dict(prod)
            cur_title = req.title.strip() if req and req.title and req.title.strip() else p_dict.get("title", b)
            cur_price = float(req.price) if req and req.price is not None else float(p_dict.get("price") or 0.0)
            label_p = float(p_dict.get("label_price") or cur_price)

        # 2. Bugün için eklenmiş bir product_history kaydı var mı kontrol et
        cursor.execute("""
            SELECT id FROM product_history
            WHERE barcode = ?
              AND substr(created_at, 1, 10) = ?
              AND sync_id IS NULL;
        """, (b, today_str))
        existing_rows = cursor.fetchall()

        if existing_rows:
            # Varsa bugün bu barkod için olan kayıtları sil (Toggle OFF)
            ids = [r["id"] for r in existing_rows]
            placeholders = ",".join("?" for _ in ids)
            cursor.execute(f"DELETE FROM product_history WHERE id IN ({placeholders});", ids)
            conn.commit()
            return success_response(
                data={"barcode": b, "is_in_changes": False},
                message=f"'{cur_title}' bugünün değişenler listesinden kaldırıldı."
            )
        else:
            # Yoksa bugünün değişenlerine ekle (Toggle ON)
            diff = round(cur_price - label_p, 2) if label_p > 0 else cur_price
            diff_pct = round((diff / label_p * 100) if label_p > 0 else (100.0 if cur_price > 0 else 0.0), 1)

            cursor.execute("""
                INSERT INTO product_history (
                    barcode, event_type, old_title, new_title,
                    old_price, new_price, diff_amount, diff_percent,
                    source, device_name, details, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (
                b, "manual_flag", cur_title, cur_title,
                label_p if label_p > 0 else 0.0, cur_price,
                diff, diff_pct,
                dev_name, dev_name,
                "Kullanıcı tarafından değişenlere manuel eklendi.",
                now_str
            ))
            conn.commit()
            return success_response(
                data={"barcode": b, "is_in_changes": True},
                message=f"'{cur_title}' bugünün değişenler listesine eklendi! 📌"
            )

@router.put("/products/{barcode}")
async def update_product(barcode: str, req: ProductUpdateRequest):
    res = update_product_details(
        barcode=barcode,
        title=req.title,
        price=req.price,
        brand=req.brand,
        unit=req.unit,
        device_name=req.device_name or "Ana PC"
    )
    if not res.get("success"):
        return error_response(message=res.get("message", "Ürün güncellenemedi."), status_code=400)
    return success_response(data=res, message=res["message"])

@router.get("/products/{barcode}/history")
async def get_product_history(barcode: str):
    history = get_full_product_history(barcode)
    prod = get_product_by_barcode(barcode)
    return success_response(
        data={
            "barcode": barcode,
            "product": prod,
            "history": history,
            "count": len(history)
        },
        message="Ürün hareket ve değişiklik geçmişi listelendi"
    )

@router.post("/products/history/{history_id}/revert")
async def revert_history(history_id: int):
    res = revert_product_history(history_id)
    if not res.get("success"):
        return error_response(message=res.get("message", "Geri alma başarısız oldu."), status_code=400)
    return success_response(data=res, message=res["message"])

@router.post("/products/sync-label-prices")
async def sync_label_prices():
    count = sync_all_label_prices_to_pos_price()
    return success_response(
        data={"updated_count": count},
        message=f"{count} ürünün etiket fiyatı güncel satış fiyatına eşitlendi."
    )

@router.post("/products/{barcode}/confirm-printed")
async def confirm_product_printed(barcode: str):
    """Fiziki etiket basıldığında etiket fiyatını kasa fiyatına eşitler ve basım tarihini günceller."""
    res = update_product_printed_time(barcode)
    return success_response(
        data=res,
        message="Ürün etiket fiyatı güncellendi ve basıldı olarak onaylandı."
    )

@router.post("/products/quick-update-clipboard")
async def quick_update_clipboard_endpoint(
    raw_text: str = Form(...),
    device_name: Optional[str] = Form("Ana PC - Fiyat Güncelleme Masası"),
    preview_only: bool = Form(False)
):
    dev_name = device_name.strip() if device_name and device_name.strip() else "Ana PC - Fiyat Güncelleme Masası"

    if not raw_text or not raw_text.strip():
        return error_response(message="Lütfen güncellenecek ürün tablosunu yapıştırın (Ctrl + V).", status_code=400)

    items, blacklisted_items = parse_raw_text_products(raw_text, collect_blacklisted=True)
    if not items and not blacklisted_items:
        return error_response(message="Yapıştırılan veriden geçerli barkod veya ürün satırı ayrıştırılamadı.", status_code=400)

    if preview_only:
        preview_data = preview_clipboard_price_update(items, blacklisted_items=blacklisted_items, device_name=dev_name)
        return success_response(data=preview_data, message=preview_data.get("message", "Önizleme hesaplandı."))

    if not items:
        return error_response(
            message=f"Güncellenecek geçerli ürün bulunamadı. Yapıştırılan {len(blacklisted_items)} ürün kara listede/engellenmiş.",
            status_code=400
        )

    res = update_products_by_clipboard_data(items, device_name=dev_name, blacklisted_count=len(blacklisted_items))
    if res.get("success"):
        return success_response(data=res, message=res.get("message", "Fiyatlar başarıyla güncellendi."))
    return error_response(message=res.get("message", "Güncelleme başarısız oldu."), status_code=400)


@router.post("/products/quick-update-excel")
async def quick_update_excel_endpoint(
    file: UploadFile = File(...),
    device_name: Optional[str] = Form("Ana PC - Excel Yükleme"),
    preview_only: bool = Form(True)
):
    dev_name = device_name.strip() if device_name and device_name.strip() else "Ana PC - Excel Yükleme"
    filename = file.filename or "fiyat.xlsx"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in ['.xlsx', '.xls', '.xlsm', '.csv', '.txt']:
        return error_response(message="Lütfen geçerli bir Excel (.xlsx, .xls) veya CSV dosyası seçin.", status_code=400)

    try:
        suffix = ext
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp_path = tmp.name
            content = await file.read()
            tmp.write(content)

        items, blacklisted_items = parse_vegawin_file(tmp_path, collect_blacklisted=True)
        try:
            os.remove(tmp_path)
        except Exception:
            pass

        if not items and not blacklisted_items:
            return error_response(message="Yüklenen Excel dosyasından geçerli ürün veya fiyat sütunları okunamadı. Lütfen dosya formatını kontrol edin.", status_code=400)

        # Tabloya aktarmak için satırları döndür (Spreadsheet 4 Temel Sütun: Barkod, Malın Cinsi, Fiyat, Değişme Tarihi)
        grid_rows = []
        # Başlık satırı
        grid_rows.append(["Barkod", "Malın Cinsi", "1. Satış Fiyatı", "1. Fiyat Değişme Tarihi"])
        for it in items:
            p_val = f"{it['price']:.2f} TL" if it.get("price") is not None else ""
            d_val = it.get("price_updated_at", "")
            grid_rows.append([
                it.get("barcode", ""),
                it.get("title", ""),
                p_val,
                d_val
            ])

        # TSV formatında metin oluştur (daha sonra doğrudan güncelleme için)
        tsv_lines = ["\t".join(r) for r in grid_rows]
        raw_text = "\n".join(tsv_lines)

        preview_data = preview_clipboard_price_update(items, blacklisted_items=blacklisted_items, device_name=dev_name)
        preview_data["grid_rows"] = grid_rows
        preview_data["raw_text"] = raw_text
        preview_data["filename"] = filename

        return success_response(
            data=preview_data,
            message=f"Excel dosyasından ({filename}) toplam {len(items)} ürün başarıyla okundu."
        )
    except Exception as e:
        return error_response(message=f"Excel dosyası işlenirken hata oluştu: {str(e)}", status_code=500)



@router.get("/export/products/excel")
async def export_products_excel(only_new: bool = False):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    products = get_all_products(limit=100000, only_new=only_new)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ürün Listesi"

    headers = ["Barkod", "Stok Kodu", "Malın Cinsi", "Fiyat (TL)", "Birim", "Yeni Ürün", "Son Güncelleme"]
    ws.append(headers)

    # Başlık stili
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for p in products:
        ws.append([
            p.get("barcode", ""),
            p.get("stock_code", ""),
            p.get("title", ""),
            float(p.get("price", 0) or 0),
            p.get("unit", "ADET"),
            "Evet" if p.get("is_new") else "Hayır",
            p.get("updated_at", "")
        ])

    # Sütun genişliklerini otomatik ayarla
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=urunler_listesi.xlsx"}
    )

@router.get("/export/products/csv")
async def export_products_csv(only_new: bool = False):
    products = get_all_products(limit=100000, only_new=only_new)
    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Barkod", "Stok Kodu", "Ürün Adı", "Fiyat (TL)", "Yeni Ürün", "Birim", "Güncellenme Tarihi"])
    for p in products:
        writer.writerow([
            p.get("barcode", ""),
            p.get("stock_code", ""),
            p.get("title", ""),
            str(p.get("price", 0)).replace('.', ','),
            "Evet" if p.get("is_new") else "Hayır",
            p.get("unit", "ADET"),
            p.get("updated_at", "")
        ])
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=urunler_listesi.csv"}
    )

@router.get("/export/products/json")
async def export_products_json(only_new: bool = False):
    products = get_all_products(limit=100000, only_new=only_new)
    return JSONResponse(
        content=products,
        headers={"Content-Disposition": "attachment; filename=urunler_listesi.json"}
    )

